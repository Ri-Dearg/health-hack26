"""Excel round trip with TUH's equipment registers, for the changeover period.

Export: a Read me tab, one tab per category, and a Waitlist tab. The columns use the
register's own headings (Asset No., Type, HC21 No., Status, ...) so staff recognise them.

Import: reads our export or TUH's registers as they are (e.g. "Asset Register" and
"Cushion Register"). Each exported row carries a hidden fingerprint of what it said at
export, so a re-import can tell who changed what:

    unchanged in Excel                  -> skipped (the app stays the source of truth)
    changed in Excel only               -> applied
    changed in Excel and in the app     -> conflict: listed for a person, nothing applied
    no fingerprint (TUH's register
    or a row added by hand)             -> applied

What the importer understands about TUH's registers (from the transcribed Wheelchair Asset
Register, 9 October 2026; anything marked "assumed" needs confirming with the OT team):
- Status codes: OWWP (assumed "on ward with patient") = on loan; STL = short-term loan;
  OT Dept / Ready for use / a storeroom = available; Workshop = awaiting repair; WASH =
  awaiting repair (the brief has no cleaning status); Missing / Awaiting reimbursement = missing.
  Anything else is flagged for a person to check. The exact text is always kept as "Status on register".
- Grey section rows ("For the provision of clients in Burkitt / ARHC") start a pool.
- Row colours: yellow Asset No. = flagged; red = missing. AT+ / AT- = on AssetTrak or not (assumed).
- "Updated" is when the row was last checked; it becomes the item's location date.

Only columns present in the sheet are applied, so a register without a "Pool" column doesn't
blank the pool. Patient names are never imported: a name column is ignored with a warning.
"""

import hashlib
import io
import json
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.models import Staff

from . import services
from .models import Category, Equipment, Loan, LocationLog, Pool, WaitlistEntry

S = Equipment.Status

# (key, header, width). The order is the export's column order: the register's own columns
# first, in its order, then the ones the app adds.
ITEM_COLUMNS = [
    ("on_assettrak", "AssetTrak", 9),
    ("asset_number", "Asset No.", 12),
    ("description", "Type", 28),
    ("size", "Size (WxDxH)", 18),
    ("serial_number", "Serial No.", 16),
    ("hc21_number", "HC21 No.", 11),
    ("dml_number", "DML No.", 11),
    ("max_weight", "Max Wgt", 13),
    ("issued_by", "Therapist", 16),
    ("mrn", "MRN", 12),
    ("location", "Ward / Location", 20),
    ("updated", "Updated", 11),
    ("status", "Status", 24),
    ("comments", "Comments", 32),
    ("last_service", "Last Service date", 13),
    ("purchase_date", "Purchased in", 12),
    ("replacement_cover", "Replacement cover", 10),
    ("category", "Category", 26),
    ("pool", "Pool", 24),
    ("barcode", "Barcode", 14),
    ("on_loan", "On loan", 8),
    ("ward", "Ward", 10),
    ("bed", "Bed", 6),
    ("issued", "Date issued", 12),
    ("due", "Due back", 12),
    ("short_term", "Short-term loan", 10),
    ("flagged", "Flagged", 8),
]
DATE_KEYS = ("updated", "issued", "due", "purchase_date", "last_service")
YES_NO_KEYS = ("on_loan", "short_term", "flagged", "replacement_cover", "on_assettrak")
WAITLIST_COLUMNS = [
    ("mrn", "MRN", 14),
    ("category", "Category", 30),
    ("ward", "Ward", 12),
    ("priority", "Priority", 20),
    ("requested", "Requested", 12),
    ("wl_status", "Status", 12),
]
FINGERPRINT_HEADER = "Row check (do not edit)"

# Header spellings we accept, after lower-casing and stripping punctuation.
# asset_number must come first: it's how the header row is found.
ITEM_ALIASES = {
    "asset_number": ["asset no", "asset number", "asset id", "asset tag"],
    "on_assettrak": ["assettrak", "asset trak", "asset"],  # the register's column A holds AT+ / AT-
    "category": ["category", "equipment category"],
    "description": ["type", "description", "make model", "model", "item"],
    "size": ["size wxdxh", "size"],
    "barcode": ["barcode", "bar code"],
    "serial_number": ["serial no", "serial number", "serial", "s n", "sn"],
    "hc21_number": ["hc21 no", "hc21 number", "hc21"],
    "dml_number": ["dml no", "dml number", "dml"],
    "max_weight": ["max wgt", "wgt capacity", "max weight", "weight capacity", "max user weight"],
    "pool": ["pool"],
    "status": ["status", "state"],
    "location": ["ward location", "location", "current location"],
    "updated": ["updated", "last updated", "date updated"],
    "on_loan": ["on loan", "loaned", "on loan y n"],
    "mrn": ["mrn", "patient mrn", "medical record number", "medical records number", "mrn number", "hospital number"],
    "ward": ["ward"],
    "bed": ["bed", "bed number", "bed no"],
    "issued_by": ["therapist", "issued by", "prescribing ot", "prescribing therapist"],
    "issued": ["date issued", "issue date", "issued", "date of issue", "date out"],
    "due": ["due back", "due date", "return by"],
    "short_term": ["short term loan", "short term", "discharged with patient", "home loan"],
    "purchase_date": ["purchased in", "purchase date", "date purchased", "purchased"],
    "last_service": ["last service date", "last service", "service date", "last serviced", "service information",
                     "service info"],
    "replacement_cover": ["replacement cover", "spare cover"],
    "flagged": ["flagged", "flag"],
    "comments": ["comments", "comment", "notes"],
    "fingerprint": ["row check do not edit"],
}
WAITLIST_ALIASES = {
    "mrn": ITEM_ALIASES["mrn"],
    "category": ITEM_ALIASES["category"] + ["equipment required", "equipment needed", "item required"],
    "ward": ["ward"],
    "priority": ["priority", "clinical priority", "clinical need"],
    "requested": ["requested", "date requested", "date added", "date"],
    "wl_status": ["status"],
}
# Never stored, whatever the spreadsheet says.
NAME_HEADERS = {"patient name", "name", "patient", "first name", "forename", "surname", "last name", "full name"}
SKIP_SHEETS = {"read me", "readme", "instructions", "legend", "key", "summary", "notes"}
WAITLIST_SHEETS = {"waitlist", "wait list", "waiting list"}
# Register tab -> category, when the sheet has no Category column. Otherwise the tab name is the category.
SHEET_CATEGORIES = {
    "asset register": "Specialised wheelchair",
    "wheelchair register": "Specialised wheelchair",
    "cushion register": "Pressure-relieving cushion",
}
BLANK_VALUES = {"n a", "na", "none", "nil", "unknown", "codes missing"}

# TUH status codes -> (status, short-term loan?, location hint). Checked in order, first match wins.
STATUS_RULES = [
    (("owwp",), S.ON_LOAN, False, None),
    (("stl", "short term", "short-term", "home loan"), S.ON_LOAN, True, None),
    (("missing", "lost", "reimbursement"), S.MISSING, False, None),
    (("workshop",), S.AWAITING_REPAIR, False, "Workshop"),
    (("repair", "service", "faulty", "broken"), S.AWAITING_REPAIR, False, None),
    (("wash", "clean", "decontam"), S.AWAITING_REPAIR, False, None),
    (("storeroom", "store rom", "store room"), S.AVAILABLE, False, "keep"),
    (("ot dept", "ot department", "ot office", "ready for use", "in store", "in stock"), S.AVAILABLE, False, "OT store"),
    (("allocated", "reserved"), S.ALLOCATED, False, None),
    (("on loan", "issued"), S.ON_LOAN, False, None),
]

# Export tints by status. Words go in the Status column too, so nobody relies on colour.
STATUS_FILLS = {
    S.AVAILABLE: "FFE2EFDA",
    S.ON_LOAN: "FFDDEBF7",
    S.AWAITING_REPAIR: "FFFCE4D6",
    S.MISSING: "FFF8CBAD",
    S.ALLOCATED: "FFE4DFEC",
}
FLAG_FILL = "FFFFFF00"
# Fill of the Asset No. cell on TUH's register: yellow = flagged, red = missing.
FLAG_FILLS = {"FFFFFF00"}
MISSING_FILLS = {"FFE02020", "FFFF0000", "FFC00000", STATUS_FILLS[S.MISSING]}


@dataclass
class ImportReport:
    dry_run: bool
    created: list = field(default_factory=list)
    updated: list = field(default_factory=list)      # (asset number, [field names])
    unchanged: int = 0
    conflicts: list = field(default_factory=list)    # (asset number, reason)
    warnings: list = field(default_factory=list)
    offers: list = field(default_factory=list)       # returned items with someone waiting
    pools: list = field(default_factory=list)        # section rows read as pools
    loans_opened: int = 0
    loans_without_patient: int = 0
    waitlist_created: int = 0
    waitlist_updated: int = 0

    @property
    def changes(self):
        return len(self.created) + len(self.updated) + self.waitlist_created + self.waitlist_updated


# ---------- normalising values so Excel and the database compare equal ----------

def _norm_header(value):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(value or "").lower())).strip()


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # numbers typed into Excel come back as floats
    if isinstance(value, (datetime, date)):
        return _date(value)
    return str(value).strip()


def _identifier(value, prefix=""):
    """Serial, HC21 and DML numbers: drop 'N/A', '?', and prefixes like 'SN:' or 'DML:'."""
    text = _text(value)
    if prefix:
        text = re.sub(rf"^{prefix}\s*:?\s*(?=\d)", "", text, flags=re.IGNORECASE)  # only when digits follow
    return "" if _norm_header(text) in BLANK_VALUES or text in {"?", "-"} else text


def _date(value):
    """Return an ISO date string, or "" if blank. Raises ValueError if it can't be read."""
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if re.fullmatch(r"\d{5}(\.0)?", text):  # an Excel date serial that ended up as text, e.g. 45931
        return (date(1899, 12, 30) + timedelta(days=int(float(text)))).isoformat()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%d.%m.%y",  # Irish day-first
                "%b-%y", "%b-%Y", "%b %Y", "%B %Y", "%m/%Y", "%Y"):                   # month or year only
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"can't read date '{text}'")


def _yes_no(value):
    """'yes', 'no' or '' (unknown). Understands AT+ / AT-, ticks and '? NO'."""
    text = _text(value).lower()
    if not text:
        return ""
    if text.endswith("+") or text.lstrip("? ").startswith(("y", "true", "1", "x", "✓", "✔")):
        return "yes"
    if text.endswith("-") or text.lstrip("? ").startswith(("n", "false", "0")):
        return "no"
    return ""


def _status(value):
    """Map register text to (status, short_term, location_hint). status is None if not understood."""
    text = _text(value).lower()
    if not text:
        return None, False, None
    for status in S:
        if text in (status.value, status.label.lower()):
            return status.value, False, None
    for words, status, short_term, hint in STATUS_RULES:
        if any(text.startswith(w) if len(w) <= 4 else w in text for w in words):
            return status.value, short_term, (_text(value) if hint == "keep" else hint)
    return None, False, None


def _priority(value):
    text = _text(value).lower()
    if "discharge" in text or "critical" in text or text == "1":
        return WaitlistEntry.Priority.DISCHARGE_CRITICAL
    if "urgent" in text or text == "2":
        return WaitlistEntry.Priority.URGENT
    return WaitlistEntry.Priority.ROUTINE


def _local_dt(iso, at=time(12)):
    return timezone.make_aware(datetime.combine(date.fromisoformat(iso), at))


def _bool_text(value):
    return "" if value is None else ("yes" if value else "no")


def db_row(item):
    """The item as canonical strings: exactly what an untouched export row parses back to."""
    loan = item.open_loan
    return {
        "on_assettrak": _bool_text(item.on_assettrak),
        "asset_number": item.asset_number,
        "description": item.description,
        "size": item.size,
        # Identifiers compare in the same cleaned-up form the importer produces ("SN: 1" = "1", "N/A" = blank).
        "serial_number": _identifier(item.serial_number, "SN"),
        "hc21_number": _identifier(item.hc21_number),
        "dml_number": _identifier(item.dml_number, "DML"),
        "max_weight": item.max_weight,
        "issued_by": (str(loan.issued_by) if loan.issued_by else loan.issued_by_name) if loan else "",
        "mrn": loan.patient.mrn if loan and loan.patient else "",
        "location": item.location,
        "updated": timezone.localdate(item.location_updated_at).isoformat(),
        "status": item.status,
        "comments": item.comments,
        "last_service": _date(item.last_service_date),
        "purchase_date": _date(item.purchase_date),
        "replacement_cover": _bool_text(item.replacement_cover),
        "category": item.category.name,
        "pool": item.pool.name,
        "barcode": _identifier(item.barcode),
        "on_loan": "yes" if loan else "no",
        "ward": loan.ward if loan else "",
        "bed": loan.bed if loan else "",
        "issued": timezone.localdate(loan.issued_at).isoformat() if loan else "",
        "due": timezone.localdate(loan.due_at).isoformat() if loan else "",
        "short_term": "yes" if loan and loan.short_term_loan else "no",
        "flagged": "yes" if item.flagged else "no",
        # Kept from TUH's register, but not exported or fingerprinted: our export's Status column is the app's.
        "register_status": item.register_status,
    }


def fingerprint(row):
    canonical = json.dumps([row.get(key, "") for key, _, _ in ITEM_COLUMNS], ensure_ascii=False)
    return hashlib.sha256(canonical.encode()).hexdigest()[:16]


# ---------- export ----------

def _sheet_title(name, used):
    title = re.sub(r"[\[\]:*?/\\]", "-", name)[:31]
    base, n = title, 2
    while title.lower() in used:
        title = f"{base[:28]} {n}"
        n += 1
    used.add(title.lower())
    return title


def _write_table(ws, columns, rows, fills=None):
    bold = Font(bold=True)
    for col, (_, header, width) in enumerate(columns, start=1):
        ws.cell(row=1, column=col, value=header).font = bold
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            if fills and fills[r - 2]:
                cell.fill = PatternFill("solid", fgColor=fills[r - 2])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def _display(row, item):
    """Canonical row -> what a person sees in Excel."""
    display = dict(row)
    display["status"] = item.get_status_display()
    display["on_assettrak"] = {"yes": "AT+", "no": "AT-"}.get(row["on_assettrak"], "")
    for key in ("on_loan", "short_term", "flagged", "replacement_cover"):
        display[key] = row[key].capitalize()
    for key in DATE_KEYS:
        display[key] = date.fromisoformat(row[key]) if row[key] else None
    return display


def build_workbook(pool=None):
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    lines = [
        ("OT equipment tracker export", None),
        ("Exported", timezone.localtime().strftime("%d/%m/%Y %H:%M")),
        ("", None),
        ("The app is the source of truth. This file is a snapshot.", None),
        ("You can edit a row and import the file back. Rows you didn't change are skipped.", None),
        ("If the app also changed a row since this export, the import lists it as a conflict and changes nothing.", None),
        ("Patient names are never stored: record the MRN only.", None),
        ("Don't edit or delete the hidden 'Row check' column.", None),
        ("A yellow Asset No. means the item is flagged for someone to check.", None),
        ("", None),
        ("Status key (the Status column always says it in words)", None),
    ]
    for r, (a, b) in enumerate(lines, start=1):
        readme.cell(row=r, column=1, value=a)
        if b:
            readme.cell(row=r, column=2, value=b)
    readme["A1"].font = Font(bold=True, size=14)
    for r, status in enumerate(S, start=len(lines) + 1):
        readme.cell(row=r, column=1, value=status.label).fill = PatternFill("solid", fgColor=STATUS_FILLS[status])
    readme.column_dimensions["A"].width = 100

    items = Equipment.objects.select_related("category", "pool").prefetch_related("loans__patient", "loans__issued_by")
    if pool:
        items = items.filter(pool__name=pool)
    used = {"read me", "waitlist"}
    columns = ITEM_COLUMNS + [("fingerprint", FINGERPRINT_HEADER, 4)]
    keys = [k for k, _, _ in ITEM_COLUMNS]
    for category in Category.objects.all():
        cat_items = [i for i in items if i.category_id == category.pk]
        if not cat_items:
            continue
        ws = wb.create_sheet(_sheet_title(category.name, used))
        rows, fills = [], []
        for item in cat_items:
            row = db_row(item)
            display = _display(row, item)
            rows.append([display[key] for key in keys] + [fingerprint(row)])
            fills.append(STATUS_FILLS[item.status])
        _write_table(ws, columns, rows, fills)
        for r, item in enumerate(cat_items, start=2):
            for key in DATE_KEYS:
                ws.cell(row=r, column=keys.index(key) + 1).number_format = "DD/MM/YYYY"
            if item.flagged:
                ws.cell(row=r, column=keys.index("asset_number") + 1).fill = PatternFill("solid", fgColor=FLAG_FILL)
        ws.column_dimensions[get_column_letter(len(columns))].hidden = True

    ws = wb.create_sheet("Waitlist")
    entries = WaitlistEntry.objects.exclude(status=WaitlistEntry.Status.CANCELLED).select_related("patient", "category")
    _write_table(ws, WAITLIST_COLUMNS, [
        [e.patient.mrn, e.category.name, e.ward, e.get_priority_display(),
         timezone.localdate(e.requested_at), e.get_status_display()]
        for e in entries
    ])
    for r in range(2, entries.count() + 2):
        ws.cell(row=r, column=5).number_format = "DD/MM/YYYY"
    return wb


def export_bytes(pool=None):
    buffer = io.BytesIO()
    build_workbook(pool).save(buffer)
    return buffer.getvalue()


# ---------- import ----------

def _find_header(ws, aliases, report):
    """Find the header row (within the first 10 rows) and map column index -> key."""
    lookup = {alias: key for key, names in aliases.items() for alias in names}
    first_key = next(iter(aliases))
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        headers = [_norm_header(v) for v in row]
        if not any(lookup.get(h) == first_key for h in headers):
            continue
        mapping, ignored = {}, []
        for i, h in enumerate(headers):
            if h in NAME_HEADERS:
                report.warnings.append(f"'{ws.title}': ignored the '{row[i]}' column. Patient names are never stored.")
            elif h in lookup and lookup[h] not in mapping.values():
                mapping[i] = lookup[h]
            elif h:
                ignored.append(str(row[i]))
        if ignored:
            report.warnings.append(f"'{ws.title}': columns not recognised, so not imported: {', '.join(ignored)}.")
        return r, mapping
    return None, {}


def _section_pool(values, mapping):
    """A grey section row such as 'For the provision of clients in Burkitt / ARHC' -> (pool name, full text)."""
    asset_col = next(i for i, k in mapping.items() if k == "asset_number")
    text = _text(values[asset_col]) if asset_col < len(values) else ""
    others = [values[i] for i in mapping if i != asset_col and i < len(values)]
    if not text or any(_text(v) for v in others) or not (len(text) >= 20 or text.lower().startswith("for ")):
        return None
    match = re.search(r"\bin (.+)$", text)
    return (match.group(1).strip() if match else text)[:100], text


def _parse_item_row(values, mapping, ws, asset_fill, report):
    from_register = "fingerprint" not in mapping.values()  # TUH's own file rather than our export
    raw = {key: values[i] if i < len(values) else None for i, key in mapping.items()}
    asset = _text(raw.get("asset_number"))
    row, short_term_from_status, hint, extra_notes = {}, False, None, []
    headers = {key: header for key, header, _ in ITEM_COLUMNS}
    for key, value in raw.items():
        try:
            if key in DATE_KEYS:
                try:
                    row[key] = _date(value)
                except ValueError:
                    row[key] = ""  # text in a date column, e.g. "Missing since ...": keep it as a note
                    extra_notes.append(f"{headers[key]}: {_text(value)}")
                    report.warnings.append(f"'{ws.title}' {asset}: '{_text(value)[:40]}' isn't a date. Kept in Comments.")
            elif key in YES_NO_KEYS:
                row[key] = _yes_no(value)
            elif key == "status":
                if from_register and _text(value):
                    row["register_status"] = _text(value)
                status, short_term_from_status, hint = _status(value)
                if status:
                    row[key] = status
                elif _text(value):
                    row["flagged"] = "yes"
                    row["_unknown_status"] = _text(value)
            elif key == "serial_number":
                row[key] = _identifier(value, "SN")
            elif key == "dml_number":
                row[key] = _identifier(value, "DML")
            elif key in ("hc21_number", "barcode"):
                row[key] = _identifier(value)
            else:
                row[key] = _text(value)
        except ValueError as exc:
            report.warnings.append(f"'{ws.title}' {asset}: {exc}, ignored.")

    if unknown := row.pop("_unknown_status", None):
        if not from_register:
            extra_notes.append(f"Status typed in Excel: {unknown}")  # no "Status on register" to keep it in
        report.warnings.append(f"'{ws.title}' {asset}: status '{unknown}' not understood. Flagged for someone to check.")
    if extra_notes:
        row["comments"] = "\n".join(filter(None, [row.get("comments", ""), *extra_notes]))
    if asset_fill in FLAG_FILLS and "flagged" not in mapping.values():
        row["flagged"] = "yes"
    if asset_fill in MISSING_FILLS and "status" not in row:
        row["status"] = S.MISSING.value
    if row.get("status") == S.MISSING and not row.get("location"):
        hint = hint or "Unknown"
    if hint and not row.get("location"):
        row["location"] = hint
    if short_term_from_status and "short_term" not in row:
        row["short_term"] = "yes"
    if row.get("mrn"):
        row["mrn"] = row["mrn"].upper()
    return row


def _apply_item(row, sheet_title, default_pool, report):
    asset = row["asset_number"]
    stored_fp = row.pop("fingerprint", "")
    item = Equipment.objects.filter(asset_number=asset).select_related("category", "pool").first()

    if item:
        current = db_row(item)
        file_view = current | row  # columns missing from the sheet count as unchanged
        if file_view == current:
            report.unchanged += 1
            return
        if stored_fp:
            if fingerprint(file_view) == stored_fp:
                report.unchanged += 1  # untouched in Excel; the app has moved on, and the app wins
                return
            if fingerprint(current) != stored_fp:
                changed = [k for k in row if row[k] != current.get(k)]
                report.conflicts.append((asset, f"changed in both Excel and the app since export ({', '.join(changed)})"))
                return

    name = row.get("category") or SHEET_CATEGORIES.get(_norm_header(sheet_title), sheet_title)
    category = Category.objects.filter(name__iexact=name).first()
    if not category:
        category = Category.objects.create(name=name)
        report.warnings.append(f"New category created: '{name}'.")
    pool_name = row.get("pool") or (item.pool.name if item else default_pool)
    pool = Pool.objects.filter(name__iexact=pool_name).first() or Pool.objects.create(name=pool_name)

    before = db_row(item) if item else {}
    if not item:
        item = Equipment(asset_number=asset)
    item.category, item.pool = category, pool
    for key in ("description", "size", "serial_number", "hc21_number", "dml_number", "max_weight", "barcode",
                "comments", "register_status"):
        if key in row:
            setattr(item, key, row[key])
    for key, attr in [("purchase_date", "purchase_date"), ("last_service", "last_service_date")]:
        if key in row:
            setattr(item, attr, date.fromisoformat(row[key]) if row[key] else None)
    for key in ("on_assettrak", "replacement_cover"):
        if key in row:
            setattr(item, key, {"yes": True, "no": False}.get(row[key]))
    if "flagged" in row:
        item.flagged = row["flagged"] == "yes"
    item.save()

    # Loan state: the Status column wins; otherwise the On loan column.
    if "status" in row:
        wants_loan = row["status"] == S.ON_LOAN
    elif "on_loan" in row:
        wants_loan = row["on_loan"] == "yes"
    else:
        wants_loan = None
    loan = item.open_loan
    staff = _match_staff(row.get("issued_by"), asset, report)
    if wants_loan and not loan:
        # The register rarely has an issue date; "Updated" is the latest date it was known to be out.
        issued = row.get("issued") or row.get("updated")
        services.issue_item(
            item, staff=staff, mrn=row.get("mrn", ""), ward=row.get("ward", ""), bed=row.get("bed", ""),
            source=LocationLog.Source.SPREADSHEET, issued_at=_local_dt(issued) if issued else None,
            due_at=_local_dt(row["due"]) if row.get("due") else None,
            short_term_loan=row.get("short_term") == "yes",
        )
        if not staff and row.get("issued_by"):
            Loan.objects.filter(pk=item.open_loan.pk).update(issued_by_name=row["issued_by"])
        report.loans_opened += 1
        if not row.get("mrn"):
            report.loans_without_patient += 1
        if not row.get("location") and not row.get("ward"):
            row["location"] = "On a ward (not recorded)"
    elif wants_loan and loan:
        if "mrn" in row:
            loan.patient = services.get_patient(row["mrn"]) if row["mrn"] else None
        for key in ("ward", "bed"):
            if key in row:
                setattr(loan, key, row[key])
        if row.get("due"):
            loan.due_at = _local_dt(row["due"])
        if "short_term" in row:
            loan.short_term_loan = row["short_term"] == "yes"
        if staff:
            loan.issued_by, loan.issued_by_name = staff, ""
        elif "issued_by" in row:
            loan.issued_by_name = row["issued_by"]
        loan.save()
    elif wants_loan is False and loan:
        offer = services.return_item(item, source=LocationLog.Source.SPREADSHEET)
        if offer:
            report.offers.append(f"{asset} is back. Next on the waitlist: {offer}.")

    item.refresh_from_db()
    if "status" in row:
        item.status = row["status"]
    if row.get("location"):
        item.location = row["location"]
    if row.get("updated"):
        item.location_updated_at = _local_dt(row["updated"])
    item.save()

    after = db_row(item)
    if not before:
        report.created.append(asset)
    else:
        changed = [k for k in after if after[k] != before.get(k)]
        if changed:
            report.updated.append((asset, changed))
        else:
            report.unchanged += 1


def _match_staff(text, asset, report):
    if not text:
        return None
    text = text.strip()
    for staff in Staff.objects.filter(is_active=True):
        if text.lower() in {staff.username.lower(), str(staff).lower(), (staff.staff_number or "").lower()}:
            return staff
    report.warnings.append(f"{asset}: Therapist '{text}' doesn't match a staff account. Kept as written; "
                           "the loan needs an owner until someone claims it.")
    return None


def _check_duplicates(rows, report):
    """Identifiers that should be unique but appear on more than one row (e.g. two chairs, one HC21 No.)."""
    for key, label in [("barcode", "barcode"), ("hc21_number", "HC21 No."), ("dml_number", "DML No."),
                       ("serial_number", "serial number")]:
        seen = defaultdict(list)
        for row in rows:
            if row.get(key):
                seen[row[key]].append(row["asset_number"])
        for value, assets in seen.items():
            if len(assets) > 1:
                report.warnings.append(f"{', '.join(assets)} share the same {label} ({value}). Check which is right.")


def _apply_waitlist(ws, report):
    header_row, mapping = _find_header(ws, WAITLIST_ALIASES, report)
    if not header_row:
        report.warnings.append(f"'{ws.title}': no 'MRN' header found, tab skipped.")
        return
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        raw = {key: values[i] if i < len(values) else None for i, key in mapping.items()}
        mrn, cat_name = _text(raw.get("mrn")).upper(), _text(raw.get("category"))
        if not mrn or not cat_name:
            continue
        category = Category.objects.filter(name__iexact=cat_name).first()
        if not category:
            report.warnings.append(f"Waitlist: category '{cat_name}' not found, row for {mrn} skipped.")
            continue
        patient = services.get_patient(mrn)
        entry = WaitlistEntry.objects.filter(patient=patient, category=category).exclude(
            status=WaitlistEntry.Status.CANCELLED).first()
        fields = {}
        if "priority" in mapping.values():
            fields["priority"] = _priority(raw.get("priority"))
        if "ward" in mapping.values():
            fields["ward"] = _text(raw.get("ward"))
        if "wl_status" in mapping.values() and _text(raw.get("wl_status")):
            text = _text(raw["wl_status"]).lower()
            fields["status"] = next((s.value for s in WaitlistEntry.Status if text in (s.value, s.label.lower())),
                                    WaitlistEntry.Status.WAITING.value)
        if raw.get("requested"):
            try:
                fields["requested_at"] = _local_dt(_date(raw["requested"]), time(9))
            except ValueError as exc:
                report.warnings.append(f"Waitlist {mrn}: {exc}, ignored.")
        if not entry:
            WaitlistEntry.objects.create(patient=patient, category=category, **fields)
            report.waitlist_created += 1
        else:
            if entry.requested_at and "requested_at" in fields and \
                    timezone.localdate(entry.requested_at) == fields["requested_at"].date():
                del fields["requested_at"]  # same day: keep the precise time
            diff = {k: v for k, v in fields.items() if getattr(entry, k) != v}
            if diff:
                for k, v in diff.items():
                    setattr(entry, k, v)
                entry.save()
                report.waitlist_updated += 1


def import_workbook(file, *, dry_run=True, default_pool=Pool.DEFAULT_NAME):
    """Read an .xlsx (path or file object). With dry_run, report what would change and change nothing."""
    report = ImportReport(dry_run=dry_run)
    wb = load_workbook(file, data_only=True)
    with transaction.atomic():
        all_rows = []
        for ws in wb.worksheets:
            title = _norm_header(ws.title)
            if title in SKIP_SHEETS:
                continue
            if title in WAITLIST_SHEETS:
                _apply_waitlist(ws, report)
                continue
            header_row, mapping = _find_header(ws, ITEM_ALIASES, report)
            if not header_row:
                report.warnings.append(f"'{ws.title}': no 'Asset No.' header found, tab skipped.")
                continue
            asset_col = next(i for i, k in mapping.items() if k == "asset_number")
            seen, section_pool = set(), None
            for row_cells in ws.iter_rows(min_row=header_row + 1):
                values = [c.value for c in row_cells]
                if section := _section_pool(values, mapping):
                    section_pool = section[0]
                    Pool.objects.get_or_create(name=section_pool, defaults={"description": section[1]})
                    report.pools.append(f"'{ws.title}': rows after \"{section[1]}\" go in pool '{section_pool}'.")
                    continue
                cell = row_cells[asset_col] if asset_col < len(row_cells) else None
                fill = cell.fill.fgColor.rgb if cell is not None and cell.fill.fill_type == "solid" else None
                row = _parse_item_row(values, mapping, ws, fill if isinstance(fill, str) else None, report)
                if not row.get("asset_number"):
                    continue
                if section_pool and "pool" not in row:
                    row["pool"] = section_pool
                if row["asset_number"] in seen:
                    report.warnings.append(f"'{ws.title}': asset {row['asset_number']} appears twice, second row skipped.")
                    continue
                seen.add(row["asset_number"])
                all_rows.append(dict(row))
                _apply_item(row, ws.title, default_pool, report)
        _check_duplicates(all_rows, report)
        no_barcode = Equipment.objects.filter(asset_number__in=[r["asset_number"] for r in all_rows], barcode="").count()
        if no_barcode:
            report.warnings.append(f"{no_barcode} item(s) have no barcode. Add one before printing labels "
                                   "(Equipment list, filter 'Has barcode: No').")
        if report.loans_without_patient:
            report.warnings.append(
                f"{report.loans_without_patient} item(s) are out on a ward with no patient or therapist recorded. "
                "They're open loans that need an owner; the issue date is taken from 'Updated'."
            )
        if dry_run:
            transaction.set_rollback(True)
    return report
