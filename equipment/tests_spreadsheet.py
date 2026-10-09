import io
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from equipment import services
from equipment.models import Category, Equipment, Loan, Patient, WaitlistEntry
from equipment.spreadsheet import ITEM_ALIASES, ImportReport, _find_header, export_bytes, import_workbook


def edit(xlsx, sheet, asset, header, value):
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[headers.index("Asset No.")].value == asset:
            row[headers.index(header)].value = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class RoundTripTests(TestCase):
    def setUp(self):
        self.boards = Category.objects.create(name="Sliding board")
        self.a = Equipment.objects.create(asset_number="SLB-001", category=self.boards, barcode="5012345")
        self.b = Equipment.objects.create(asset_number="SLB-002", category=self.boards)
        services.issue_item(self.b, mrn="MRN9", ward="Ward 2B")
        WaitlistEntry.objects.create(category=self.boards, patient=services.get_patient("MRN7"),
                                     priority=WaitlistEntry.Priority.URGENT)

    def test_untouched_export_imports_with_no_changes(self):
        report = import_workbook(io.BytesIO(export_bytes()), dry_run=False)
        self.assertEqual(report.changes, 0, report)
        self.assertEqual(report.conflicts, [])
        self.assertEqual(report.unchanged, 2)

    def test_edit_in_excel_is_applied(self):
        xlsx = edit(export_bytes(), "Sliding board", "SLB-001", "Comments", "Strap frayed")
        report = import_workbook(io.BytesIO(xlsx), dry_run=False)
        self.assertEqual(report.updated, [("SLB-001", ["comments"])])
        self.a.refresh_from_db()
        self.assertEqual(self.a.comments, "Strap frayed")

    def test_dry_run_changes_nothing(self):
        xlsx = edit(export_bytes(), "Sliding board", "SLB-001", "Comments", "Strap frayed")
        self.assertEqual(len(import_workbook(io.BytesIO(xlsx), dry_run=True).updated), 1)
        self.a.refresh_from_db()
        self.assertEqual(self.a.comments, "")

    def test_change_in_app_only_is_kept(self):
        xlsx = export_bytes()
        Equipment.objects.filter(pk=self.a.pk).update(status=Equipment.Status.MISSING)
        report = import_workbook(io.BytesIO(xlsx), dry_run=False)
        self.assertEqual(report.changes, 0)
        self.a.refresh_from_db()
        self.assertEqual(self.a.status, Equipment.Status.MISSING)

    def test_change_in_both_is_a_conflict_and_not_applied(self):
        xlsx = edit(export_bytes(), "Sliding board", "SLB-001", "Comments", "Excel edit")
        Equipment.objects.filter(pk=self.a.pk).update(comments="App edit")
        report = import_workbook(io.BytesIO(xlsx), dry_run=False)
        self.assertEqual([c[0] for c in report.conflicts], ["SLB-001"])
        self.a.refresh_from_db()
        self.assertEqual(self.a.comments, "App edit")

    def test_marking_returned_in_excel_closes_loan_and_offers_waitlist(self):
        xlsx = edit(export_bytes(), "Sliding board", "SLB-002", "Status", "Available")
        xlsx = edit(xlsx, "Sliding board", "SLB-002", "On loan", "No")
        report = import_workbook(io.BytesIO(xlsx), dry_run=False)
        self.assertFalse(Loan.objects.open().exists())
        self.assertIn("MRN7", report.offers[0])

    def test_waitlist_round_trip(self):
        report = import_workbook(io.BytesIO(export_bytes()), dry_run=False)
        self.assertEqual((report.waitlist_created, report.waitlist_updated), (0, 0))


class OriginalTrackerTests(TestCase):
    """A made-up file shaped like the brief's description of TUH's tracker."""

    def make_tracker(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sliding boards"
        ws.append(["Asset No.", "Barcode", "Serial No", "Patient Name", "MRN", "Location", "Date Issued", "Service info"])
        ws.append(["SB-01", 5012345678900, "S1", "Invented Person", "mrn123", "Ward 3C", "05/10/2026", "01/03/2026"])
        ws.append(["SB-02", None, None, None, None, "OT store", None, None])
        for cell in ws[3]:
            cell.fill = PatternFill("solid", fgColor="FFFF0000")  # red: missing (assumed key)
        wl = wb.create_sheet("Waitlist")
        wl.append(["MRN", "Equipment required", "Ward", "Priority"])
        wl.append(["MRN555", "Sliding boards", "AMU", "Discharge critical"])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def test_first_import_of_tuh_tracker(self):
        report = import_workbook(io.BytesIO(self.make_tracker()), dry_run=False)
        self.assertEqual(sorted(report.created), ["SB-01", "SB-02"])
        self.assertTrue(any("Patient names are never stored" in w for w in report.warnings))

        sb1 = Equipment.objects.get(asset_number="SB-01")
        self.assertEqual(sb1.category.name, "Sliding boards")
        self.assertEqual(sb1.barcode, "5012345678900")
        self.assertEqual(str(sb1.last_service_date), "2026-03-01")
        loan = sb1.open_loan
        self.assertIsNone(loan, "no Status or On loan column, so no loan is guessed")
        self.assertEqual(Equipment.objects.get(asset_number="SB-02").status, Equipment.Status.MISSING)
        self.assertEqual(list(Patient.objects.values_list("mrn", flat=True)), ["MRN555"])
        self.assertEqual(WaitlistEntry.objects.get().priority, WaitlistEntry.Priority.DISCHARGE_CRITICAL)

    def test_on_loan_column_opens_loan_with_issue_date(self):
        wb = load_workbook(io.BytesIO(self.make_tracker()))
        ws = wb["Sliding boards"]
        ws["I1"], ws["I2"] = "On loan", "Y"
        out = io.BytesIO()
        wb.save(out)
        import_workbook(io.BytesIO(out.getvalue()), dry_run=False)
        loan = Equipment.objects.get(asset_number="SB-01").open_loan
        self.assertEqual(loan.patient.mrn, "MRN123")
        self.assertEqual(timezone.localdate(loan.issued_at).isoformat(), "2026-10-05")
        self.assertEqual(timezone.localdate(loan.due_at).isoformat(), "2026-11-04")  # 30 days, across the clock change
        self.assertTrue(loan.needs_owner)


class TuhRegisterTests(TestCase):
    """Invented rows in the shape of TUH's transcribed Wheelchair Asset Register (9 October 2026)."""

    HEADERS = ["Asset", "Asset No.", "Type", "Size (WxDxH)", "Serial No.", "HC21 No.", "DML No.", "Max Wgt",
               "Therapist", "Patient Name", "MRN", "Ward / Location", "Updated", "Status", "Comments",
               "Last Service date", "Purchased in"]

    def make_register(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Asset Register"
        ws.append(self.HEADERS)
        rows = [
            ["AT+", "X01", "Test Chair", '17"', "SN: 111", "9000001", "DML500", "135kg", None, None, None, None,
             "01/09/2026", "OWWP", None, "03.07.2025", None],
            ["AT-", "X02", "Test Chair", '18"', "222", "9000002", "500002", "135kg", None, None, None, None,
             "29/09/2026", "OT Dept", None, "29/09/2026", "May-26"],
            ["AT-", "X03", "Test Chair", '18"', "333", "9000003", "500003", "135kg", None, None, None, None,
             "08/10/2026", "Workshop", None, None, "45931"],
            ["AT+", "X04", "Test Chair", '15"', "444", "9000003", "500004", None, None, None, None, None,
             "19/09/2025", None, None, "Missing since 19/09/2025", None],
            ["AT+", "X05", "Test Cushion", '18x18"', "N/A", "9000005", None, "130 kg", None, None, None, None,
             "06/10/2026", "Ask Pat", None, None, "2024"],
            [None, "For the provision of clients in Test Ward", None, None, None, None, None, None, None, None,
             None, None, None, None, None, None, None],
            ["AT-", "X06", "Test Chair", '17"', "666", "9000006", "500006", "160kg", None, None, None, None,
             "05/10/2026", "STL - due back in October", None, None, None],
        ]
        for values in rows:
            ws.append(values)
        red, yellow = PatternFill("solid", fgColor="FFE02020"), PatternFill("solid", fgColor="FFFFFF00")
        for cell in ws[5]:
            cell.fill = red      # X04: red row = missing
        ws["B6"].fill = yellow  # X05: yellow Asset No. = flagged
        notes = wb.create_sheet("Notes")
        notes.append(["Transcription notes: not data"])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def test_import_reads_tuh_columns_codes_colours_and_sections(self):
        report = import_workbook(io.BytesIO(self.make_register()), dry_run=False)
        self.assertEqual(sorted(report.created), ["X01", "X02", "X03", "X04", "X05", "X06"])
        get = lambda a: Equipment.objects.select_related("category", "pool").get(asset_number=a)

        x1 = get("X01")
        self.assertEqual(x1.category.name, "Specialised wheelchair")
        self.assertEqual((x1.description, x1.size, x1.max_weight), ("Test Chair", '17"', "135kg"))
        self.assertEqual((x1.serial_number, x1.hc21_number, x1.dml_number), ("111", "9000001", "500"))
        self.assertTrue(x1.on_assettrak)
        self.assertEqual((x1.status, x1.register_status), (Equipment.Status.ON_LOAN, "OWWP"))
        self.assertEqual(timezone.localdate(x1.open_loan.issued_at).isoformat(), "2026-09-01")  # from Updated
        self.assertTrue(x1.open_loan.needs_owner)
        self.assertEqual(str(x1.last_service_date), "2025-07-03")

        x2 = get("X02")
        self.assertEqual((x2.status, x2.location, str(x2.purchase_date)), ("available", "OT store", "2026-05-01"))
        self.assertFalse(x2.on_assettrak)
        x3 = get("X03")
        self.assertEqual((x3.status, x3.location, str(x3.purchase_date)), ("awaiting_repair", "Workshop", "2025-10-01"))
        x4 = get("X04")
        self.assertEqual((x4.status, x4.location), ("missing", "Unknown"))
        self.assertIn("Missing since 19/09/2025", x4.comments)
        x5 = get("X05")
        self.assertTrue(x5.flagged)
        self.assertEqual(x5.register_status, "Ask Pat")
        self.assertEqual((x5.serial_number, str(x5.purchase_date)), ("", "2024-01-01"))
        x6 = get("X06")
        self.assertEqual(x6.pool.name, "Test Ward")
        self.assertTrue(x6.open_loan.short_term_loan)

        self.assertFalse(Patient.objects.exists(), "no MRNs in the register, so no patients")
        self.assertTrue(any("X03, X04 share the same HC21 No." in w for w in report.warnings))
        self.assertTrue(any("Patient names are never stored" in w for w in report.warnings))

    def test_register_round_trips_through_our_export(self):
        import_workbook(io.BytesIO(self.make_register()), dry_run=False)
        report = import_workbook(io.BytesIO(export_bytes()), dry_run=True)
        self.assertEqual((report.changes, len(report.conflicts), report.unchanged), (0, 0, 6))

    def test_reimporting_the_same_register_changes_nothing(self):
        import_workbook(io.BytesIO(self.make_register()), dry_run=False)
        report = import_workbook(io.BytesIO(self.make_register()), dry_run=True)
        self.assertEqual((report.changes, report.loans_opened), (0, 0))

    def test_every_register_column_is_recognised(self):
        cushion_headers = ["AssetTrak", "Asset No.", "Type", "Size (WxDxH)", "HC21 No.", "DML No.", "Serial no",
                           "Wgt Capacity", "Therapist", "Patient Name", "MRN", "Ward/Location", "Updated", "Status",
                           "purchased in:", "Replacement cover", "Comments"]
        for headers in (self.HEADERS, cushion_headers):
            ws = Workbook().active
            ws.append(headers)
            report = ImportReport(dry_run=True)
            _, mapping = _find_header(ws, ITEM_ALIASES, report)
            self.assertEqual([h for i, h in enumerate(headers) if i not in mapping], ["Patient Name"])  # MRN only
            self.assertFalse([w for w in report.warnings if "not recognised" in w])

    def test_unmatched_therapist_is_kept_as_written(self):
        wb = load_workbook(io.BytesIO(self.make_register()))
        wb["Asset Register"]["I2"] = "Invented Therapist"  # X01, out on a ward
        out = io.BytesIO()
        wb.save(out)
        report = import_workbook(io.BytesIO(out.getvalue()), dry_run=False)
        loan = Equipment.objects.get(asset_number="X01").open_loan
        self.assertEqual((loan.issued_by, loan.issued_by_name), (None, "Invented Therapist"))
        self.assertTrue(any("Kept as written" in w for w in report.warnings))
        again = import_workbook(io.BytesIO(export_bytes()), dry_run=True)
        self.assertEqual((again.changes, len(again.conflicts)), (0, 0))
